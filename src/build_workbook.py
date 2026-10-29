from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.comments import Comment

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_IN = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_TITLE = Font(name=FONT, size=14, bold=True, color="1F4E3D")
F_SUB = Font(name=FONT, size=10, italic=True, color="595959")
FILL_HDR = PatternFill("solid", fgColor="1F4E3D")
FILL_SEC = PatternFill("solid", fgColor="E2EFDA")
FILL_IN = PatternFill("solid", fgColor="FFF2CC")
FILL_TOT = PatternFill("solid", fgColor="D9D9D9")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

N0 = '#,##0;(#,##0);"-"'
N1 = '#,##0.0;(#,##0.0);"-"'
N3 = '#,##0.000;(#,##0.000);"-"'
P1 = '0.0%;(0.0%);"-"'

wb = Workbook()


def cell(ws, ref, value, font=F_BASE, fmt=None, fill=None, bold=False, wrap=False, border=True):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, size=font.size, bold=bold or font.bold, italic=font.italic, color=font.color)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if wrap:
        c.alignment = WRAP
    if border:
        c.border = BOX
    return c


def inp(ws, ref, value, fmt=None):
    return cell(ws, ref, value, F_IN, fmt, FILL_IN)


def header(ws, row, labels, col=1):
    for i, h in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=h)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = BOX
    ws.row_dimensions[row].height = 30


def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = F_TITLE
    if sub:
        ws["A2"] = sub
        ws["A2"].font = F_SUB
    ws.sheet_view.showGridLines = False


def widths(ws, ws_widths):
    for k, v in ws_widths.items():
        ws.column_dimensions[k].width = v


def section(ws, row, text, ncols):
    for c in range(1, ncols + 1):
        ws.cell(row=row, column=c).fill = FILL_SEC
    ws.cell(row=row, column=1, value=text).font = F_BOLD


def ef(ref):
    return f"VLOOKUP({ref},Emission_Factors!$A:$D,4,FALSE)"


# =====================================================================
# 0. README
# =====================================================================
rd = wb.active
rd.title = "README"
title(rd, "CircularFibre Group (fictional) – LCA / PCF, CO₂ Accounting & Sustainable Procurement Toolkit",
      "Portfolio model: waste paper recycling → recycled cartonboard → folding cartons & gypsum board liners. All data illustrative.")
widths(rd, {"A": 28, "B": 110})
rows = [
    ("PURPOSE", "Integrated Excel toolkit covering the core tasks of a Sustainable Procurement / Carbon Accounting role: product carbon footprints (PCF/LCA), corporate GHG inventory (Scope 1-2-3), supplier CO₂ data collection & validation, ESG supplier scorecards, KPIs, decarbonisation levers, supplier survey and training plan."),
    ("BUSINESS MODEL", "Circular value chain: (1) recovered paper (RCP) is collected and sorted, (2) pulped and converted into 100% recycled cartonboard in own board mills, (3) board is converted into folding cartons in own converting plants or sold externally, (4) part of the board is used as liner paper in own gypsum board plants. Converting trim waste (~12%) is returned to the mill as RCP (closed loop)."),
    ("STANDARDS APPLIED", "ISO 14040/14044 (LCA framework), ISO 14067 (PCF), GHG Protocol Corporate Standard + Scope 3 Standard, GHG Protocol Product Standard, WBCSD PACT Pathfinder Framework v3 (data quality ratings, PCF exchange), CEPI Carbon Footprint Framework for paper & board, EN 15804 (construction products – gypsum board EPD), EcoVadis methodology (4 themes)."),
    ("SYSTEM BOUNDARY", "PCFs: cradle-to-gate (A1-A3), fossil GHG in kg CO₂e, GWP100 (IPCC AR6). Recovered paper enters burden-free (cut-off / recycled-content approach) – only collection, sorting and transport are allocated. Biogenic carbon stored in the product is reported separately as a memo item (not netted). Folding carton additionally shows an optional cradle-to-customer value."),
    ("SHEET MAP", "Company_Inputs → production volumes & activity data | Emission_Factors → EF library with sources | PCF_Board / PCF_FoldingCarton / PCF_GypsumBoard → product footprints | Supplier_Register → supplier volumes, supplier PCFs, PACT data quality, primary-vs-secondary logic | GHG_Inventory → Scope 1/2/3 by category | Supplier_Scorecard → weighted ESG scoring & risk prioritisation | Decarbonisation → levers, reduction potential, abatement cost | KPI_Dashboard → procurement & climate KPIs with charts | Supplier_Survey → questionnaire template | Training_Plan → internal awareness programme"),
    ("COLOUR LEGEND", "Blue text on yellow fill = input / assumption (edit these) · Black = formula (do not overwrite) · Green = link to another sheet · Grey rows = totals."),
    ("HOW TO UPDATE", "1) Update production volumes in Company_Inputs. 2) Replace illustrative factors in Emission_Factors with licensed data (ecoinvent, GaBi/Sphera, UBA, DESNZ, AIB residual mix). 3) Enter supplier volumes, supplier-specific PCFs and the 5 PACT data-quality indicators in Supplier_Register – the model automatically uses supplier data if DQR ≤ threshold, otherwise falls back to secondary data. 4) Enter ESG assessment results (EcoVadis, survey) in Supplier_Scorecard. 5) Review KPI_Dashboard and Decarbonisation."),
    ("DATA QUALITY LOGIC", "PACT Data Quality Rating (DQR) = average of technological, geographical, temporal representativeness, completeness and reliability (1 = very good … 5 = very poor). Supplier PCFs with DQR above the acceptance threshold are flagged and replaced by secondary EFs until the supplier improves its data."),
    ("LIMITATIONS", "Company, suppliers and all numbers are FICTIONAL and for demonstration only. Emission factors are rounded, indicative values typical of published sources – they are not verified and must be replaced before any external use (CSRD/ESRS E1, customer PCF requests, EcoVadis)."),
]
r = 4
for k, v in rows:
    cell(rd, f"A{r}", k, bold=True, wrap=True)
    cell(rd, f"B{r}", v, wrap=True)
    rd.row_dimensions[r].height = max(30, 15 * (len(v) // 100 + 1))
    r += 1

# =====================================================================
# 1. Emission factors
# =====================================================================
efs = wb.create_sheet("Emission_Factors")
title(efs, "Emission Factor Library (kg CO₂e per unit, GWP100 AR6)", "Indicative values – replace with licensed / verified data before external reporting.")
widths(efs, {"A": 8, "B": 44, "C": 10, "D": 12, "E": 18, "F": 46, "G": 40})
header(efs, 4, ["EF ID", "Activity / material", "Unit", "kg CO₂e / unit", "Scope / category", "Source (indicative)", "Notes"])
EF = [
    ("EF01", "Natural gas – combustion", "kWh", 0.201, "Scope 1", "UBA CO₂ emission factors (Hi basis)", "Combustion only; upstream in EF02"),
    ("EF02", "Natural gas – upstream (WTT)", "kWh", 0.030, "Scope 3 Cat 3", "DESNZ 2024 WTT factors (approx.)", ""),
    ("EF03", "Grid electricity DE – location-based", "kWh", 0.363, "Scope 2 LB", "UBA electricity mix DE (provisional)", ""),
    ("EF04", "Electricity – upstream & T&D losses", "kWh", 0.045, "Scope 3 Cat 3", "DESNZ 2024 WTT + T&D (approx.)", ""),
    ("EF05", "Renewable electricity (GoO / PPA)", "kWh", 0.000, "Scope 2 MB", "GHG Protocol Scope 2 Guidance", "Market-based, contractual instrument"),
    ("EF06", "Road freight – articulated HGV (avg laden)", "tkm", 0.095, "Scope 3 Cat 4/9", "DESNZ 2024 freighting goods", ""),
    ("EF07", "Rail freight (electric/diesel mix)", "tkm", 0.028, "Scope 3 Cat 4", "DESNZ 2024 / GLEC Framework", ""),
    ("EF08", "Recovered paper – collection & sorting", "kg", 0.015, "Scope 3 Cat 1", "CEPI framework, cut-off approach (approx.)", "RCP burden-free; only collection/sorting/baling"),
    ("EF09", "Native starch (maize/wheat)", "kg", 0.850, "Scope 3 Cat 1", "ecoinvent-type average (approx.)", ""),
    ("EF10", "Calcium carbonate filler (GCC)", "kg", 0.080, "Scope 3 Cat 1", "ecoinvent-type average (approx.)", ""),
    ("EF11", "Process chemicals – generic", "kg", 2.000, "Scope 3 Cat 1", "Conservative proxy", "Sizing, retention, biocides"),
    ("EF12", "Offset printing ink", "kg", 3.200, "Scope 3 Cat 1", "Industry average (approx.)", ""),
    ("EF13", "Overprint varnish", "kg", 2.500, "Scope 3 Cat 1", "Industry average (approx.)", ""),
    ("EF14", "Dispersion adhesive", "kg", 1.800, "Scope 3 Cat 1", "Industry average (approx.)", ""),
    ("EF15", "LDPE stretch film", "kg", 2.600, "Scope 3 Cat 1", "PlasticsEurope eco-profile (approx.)", ""),
    ("EF16", "Natural gypsum (quarried)", "kg", 0.004, "Scope 3 Cat 1", "EPD-type average (approx.)", ""),
    ("EF17", "FGD gypsum (by-product of coal power)", "kg", 0.002, "Scope 3 Cat 1", "By-product, economic allocation ≈ 0", "Supply declining with coal phase-out"),
    ("EF18", "Recycled gypsum (processing)", "kg", 0.008, "Scope 3 Cat 1", "EPD-type average (approx.)", "Avoids landfill & primary extraction"),
    ("EF19", "Glass fibre", "kg", 1.600, "Scope 3 Cat 1", "ecoinvent-type average (approx.)", ""),
    ("EF20", "Gypsum board additives / foaming agent", "kg", 2.300, "Scope 3 Cat 1", "Conservative proxy", ""),
    ("EF21", "Water supply & wastewater treatment", "m³", 0.300, "Scope 3 Cat 5", "Utility average (approx.)", ""),
    ("EF22", "Rejects to incineration (fossil share)", "kg", 0.600, "Scope 3 Cat 5", "Waste incineration, plastics fraction", "Pulper rejects: plastics, metals"),
    ("EF23", "Indirect procurement – spend-based", "EUR", 0.250, "Scope 3 Cat 1", "EEIO (EXIOBASE-type) average", "Use only as fallback"),
    ("EF24", "Capital goods – spend-based", "EUR", 0.350, "Scope 3 Cat 2", "EEIO (EXIOBASE-type) average", ""),
    ("EF25", "Business travel – average (air/rail/car)", "pkm", 0.150, "Scope 3 Cat 6", "DESNZ 2024 (blended)", ""),
    ("EF26", "Employee commuting – average car", "pkm", 0.170, "Scope 3 Cat 7", "DESNZ 2024", ""),
    ("EF27", "Paperboard end-of-life (fossil, DE mix)", "kg", 0.030, "Scope 3 Cat 12", "Recycling/incineration mix (approx.)", "Biogenic CO₂ excluded"),
    ("EF28", "Gypsum board end-of-life (landfill/recycling)", "kg", 0.010, "Scope 3 Cat 12", "EN 15804 C-modules (approx.)", ""),
    ("EF29", "Diesel – combustion (mobile)", "l", 2.650, "Scope 1", "UBA / DESNZ 2024", "Forklifts, yard vehicles"),
    ("EF30", "Diesel – upstream (WTT)", "l", 0.620, "Scope 3 Cat 3", "DESNZ 2024 WTT", ""),
    ("EF31", "Residual mix DE – market-based", "kWh", 0.620, "Scope 2 MB", "AIB European Residual Mix (approx.)", ""),
    ("EF32", "Biomethane – combustion + upstream (fossil)", "kWh", 0.030, "Scope 1", "Indicative; biogenic CO₂ excluded", "Used in Decarbonisation"),
]
for i, row in enumerate(EF):
    rr = 5 + i
    for j, v in enumerate(row):
        col = "ABCDEFG"[j]
        if j == 3:
            inp(efs, f"{col}{rr}", v, "0.000")
        else:
            cell(efs, f"{col}{rr}", v, wrap=(j in (5, 6)))
efs.freeze_panes = "A5"
EF_LAST = 4 + len(EF)

# =====================================================================
# 2. Company inputs
# =====================================================================
ci = wb.create_sheet("Company_Inputs")
title(ci, "Company Inputs – Reporting Year Activity Data", "Edit blue cells. Volumes are fictional but internally consistent (mass balance).")
widths(ci, {"A": 54, "B": 16, "C": 14, "D": 60})
header(ci, 4, ["Parameter", "Value", "Unit", "Source / comment"])
CI = {}


def ci_row(r, key, label, val, unit, note, is_input=True, fmt=N0):
    cell(ci, f"A{r}", label)
    if is_input:
        inp(ci, f"B{r}", val, fmt)
    else:
        cell(ci, f"B{r}", val, fmt=fmt)
    cell(ci, f"C{r}", unit)
    cell(ci, f"D{r}", note, wrap=True)
    CI[key] = f"Company_Inputs!$B${r}"


section(ci, 5, "Production", 4)
ci_row(6, "year", "Reporting year", "2025", "", "Text format", fmt="@")
ci_row(7, "board_ext", "Cartonboard sold to external converters", 260000, "t", "Sales statistics (fictional)")
ci_row(8, "fc", "Folding carton output", 105000, "t", "Converting plants output (fictional)")
ci_row(9, "gyp", "Gypsum board output", 45000, "1,000 m²", "i.e. 45 million m² (fictional)")
ci_row(10, "gyp_w", "Gypsum board mass (12.5 mm)", 8.5, "kg/m²", "Typical standard board; product data sheet", fmt=N1)
ci_row(11, "board_fc", "Board consumed by own folding carton plants", "=B8*PCF_FoldingCarton!$C$13/1000", "t", "FC output × board input per t (PCF_FoldingCarton)", False)
ci_row(12, "board_gyp", "Liner board consumed by own gypsum board plants", "=B9*PCF_GypsumBoard!$C$17/1000", "t", "Gypsum output × liner kg per 1,000 m²", False)
ci_row(13, "board", "Total cartonboard production", "=B7+B11+B12", "t", "Mass balance", False)
section(ci, 14, "Recovered paper mass balance (circularity)", 4)
ci_row(15, "rcp_gross", "Gross recovered paper input", "=B13*PCF_Board!$C$6", "t", "Board production × RCP input ratio", False)
ci_row(16, "trim", "Internal trim returned from FC plants (closed loop)", "=B11-B8", "t", "Board input − FC output", False)
ci_row(17, "rcp_ext", "External recovered paper to be purchased", "=B15-B16", "t", "Must reconcile with Supplier_Register", False)
ci_row(18, "gyp_raw", "Raw gypsum input (all sources)", "=B9*PCF_GypsumBoard!$C$6/1000", "t", "Gypsum output × gypsum kg per 1,000 m²", False)
section(ci, 19, "Other activity data", 4)
ci_row(20, "diesel", "Diesel – forklifts & yard vehicles", 1500000, "l", "Fuel cards / tank records")
ci_row(21, "green", "Share of electricity covered by GoO / PPA", 0.25, "%", "Energy procurement contracts", fmt=P1)
ci_row(22, "spend_ind", "Indirect procurement spend (services, MRO, IT)", 60000000, "EUR", "ERP spend cube, excl. direct materials")
ci_row(23, "capex", "Capital expenditure", 35000000, "EUR", "Fixed asset additions")
ci_row(24, "travel", "Business travel", 4200000, "pkm", "Travel agency report")
ci_row(25, "fte", "Employees", 2800, "FTE", "HR")
ci_row(26, "commute_km", "Average one-way commute", 18, "km", "Employee mobility survey")
ci_row(27, "workdays", "Working days per year", 210, "days", "HR")
ci_row(28, "dist_cust", "Avg distance – external board collected by customers", 200, "km", "Logistics (Cat 9: customer-arranged)")
ci_row(29, "dist_out", "Avg distance – own outbound deliveries (FC & gypsum board)", 300, "km", "Logistics (Cat 4: paid by company)")

# =====================================================================
# PCF helper
# =====================================================================
PCF_HDR = ["Stage", "Input / process", "Quantity per FU", "Unit", "EF ID", "EF (kg CO₂e / unit)", "kg CO₂e per FU", "Share", "Company scope", "Data source / quality note"]


def pcf_sheet(name, ttl, sub, params, lines, total_label):
    ws = wb.create_sheet(name)
    title(ws, ttl, sub)
    widths(ws, {"A": 22, "B": 46, "C": 14, "D": 9, "E": 8, "F": 13, "G": 14, "H": 9, "I": 16, "J": 44})
    cell(ws, "A4", "Parameters", bold=True, border=False)
    keys = {}
    r = 5
    for key, label, val, unit, note in params:
        cell(ws, f"B{r}", label)
        if isinstance(val, str) and val.startswith("="):
            cell(ws, f"C{r}", val, fmt=N3 if "ratio" in key else N0, font=F_LINK if "!" in val else F_BASE)
        else:
            inp(ws, f"C{r}", val, N3 if isinstance(val, float) and val < 10 else N0)
        cell(ws, f"D{r}", unit)
        cell(ws, f"J{r}", note, wrap=True)
        keys[key] = f"$C${r}"
        r += 1
    r += 1
    hdr = r
    header(ws, hdr, PCF_HDR)
    r += 1
    first = r
    for ln in lines:
        key, stage, label, qty, unit, efid, scope, note = ln
        cell(ws, f"A{r}", stage)
        cell(ws, f"B{r}", label)
        if isinstance(qty, str) and qty.startswith("="):
            q = qty
            for k, v in keys.items():
                q = q.replace("{" + k + "}", v)
            cell(ws, f"C{r}", q, fmt=N1)
        else:
            inp(ws, f"C{r}", qty, N1)
        cell(ws, f"D{r}", unit)
        if efid.startswith("="):
            cell(ws, f"E{r}", "link")
            cell(ws, f"F{r}", efid, F_LINK, "0.000")
        else:
            cell(ws, f"E{r}", efid)
            cell(ws, f"F{r}", f"={ef(f'E{r}')}", F_LINK, "0.000")
        cell(ws, f"G{r}", f"=C{r}*F{r}", fmt=N1)
        cell(ws, f"I{r}", scope)
        cell(ws, f"J{r}", note, wrap=True)
        keys[key] = f"$C${r}"
        keys[key + "_g"] = f"$G${r}"
        r += 1
    last = r - 1
    tot = r
    for c in "ABCDEFGHIJ":
        ws[f"{c}{tot}"].fill = FILL_TOT
        ws[f"{c}{tot}"].border = BOX
    cell(ws, f"A{tot}", total_label, bold=True, fill=FILL_TOT)
    cell(ws, f"G{tot}", f"=SUM(G{first}:G{last})", fmt=N1, bold=True, fill=FILL_TOT)
    cell(ws, f"H{tot}", f"=SUM(H{first}:H{last})", fmt=P1, bold=True, fill=FILL_TOT)
    for rr in range(first, last + 1):
        cell(ws, f"H{rr}", f"=IFERROR(G{rr}/$G${tot},0)", fmt=P1)
    keys["total"] = f"$G${tot}"
    keys["first"], keys["last"], keys["tot"] = first, last, tot
    ws.freeze_panes = ws[f"A{hdr+1}"]
    return ws, keys


def stage_summary(ws, keys, stages, start):
    cell(ws, f"A{start}", "Hotspot summary by life-cycle stage", bold=True, border=False)
    header(ws, start + 1, ["Stage", "kg CO₂e", "Share"])
    f, l, t = keys["first"], keys["last"], keys["tot"]
    for i, s in enumerate(stages):
        rr = start + 2 + i
        cell(ws, f"A{rr}", s)
        cell(ws, f"B{rr}", f"=SUMIF($A${f}:$A${l},A{rr},$G${f}:$G${l})", fmt=N1)
        cell(ws, f"C{rr}", f"=IFERROR(B{rr}/$G${t},0)", fmt=P1)
    end = start + 1 + len(stages)
    pie = PieChart()
    pie.title = "Hotspots by stage"
    pie.add_data(Reference(ws, min_col=2, min_row=start + 1, max_row=end), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=start + 2, max_row=end))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.height, pie.width = 7.5, 12
    ws.add_chart(pie, f"E{start}")
    return end


# =====================================================================
# 3. PCF Board  (FU = 1 t recycled cartonboard)
# =====================================================================
pb, PB = pcf_sheet(
    "PCF_Board",
    "PCF – 100% Recycled Cartonboard (FU: 1 t, cradle-to-gate)",
    "ISO 14067 / CEPI framework · cut-off approach for recovered paper · fossil GHG only, biogenic carbon as memo item",
    [
        ("rcp_ratio", "Recovered paper input ratio (t RCP / t board)", 1.15, "t/t", "Mill yield incl. moisture & rejects; production data"),
        ("rcp_km", "Avg RCP transport distance to mill", 150, "km", "Weighted from Supplier_Register logistics data"),
        ("bio_c", "Biogenic carbon content of board (kg C / kg)", 0.40, "kg/kg", "Typical fibre carbon content (memo only)"),
    ],
    [
        ("rcp", "A1 Raw materials", "Recovered paper (collection & sorting)", "=1000*{rcp_ratio}", "kg", "EF08", "Scope 3 Cat 1", "Burden-free fibre; secondary EF – request supplier data"),
        ("starch", "A1 Raw materials", "Native starch (surface & wet-end)", 40, "kg", "EF09", "Scope 3 Cat 1", "Supplier PCF available (see register)"),
        ("caco3", "A1 Raw materials", "Calcium carbonate filler", 30, "kg", "EF10", "Scope 3 Cat 1", ""),
        ("chem", "A1 Raw materials", "Process chemicals", 10, "kg", "EF11", "Scope 3 Cat 1", "Proxy – data gap"),
        ("rcp_tr", "A2 Transport", "Inbound transport recovered paper", "={rcp_ratio}*{rcp_km}", "tkm", "EF06", "Scope 3 Cat 4", "t × km"),
        ("gas", "A3 Manufacturing", "Natural gas – steam & drying (CHP/boiler)", 1700, "kWh", "EF01", "Scope 1", "Metered, mill energy management (ISO 50001)"),
        ("gas_up", "A3 Manufacturing", "Natural gas – upstream", "={gas}", "kWh", "EF02", "Scope 3 Cat 3", ""),
        ("el", "A3 Manufacturing", "Grid electricity (location-based)", 450, "kWh", "EF03", "Scope 2", "Metered; market-based see GHG_Inventory"),
        ("el_up", "A3 Manufacturing", "Electricity – upstream & losses", "={el}", "kWh", "EF04", "Scope 3 Cat 3", ""),
        ("water", "A3 Manufacturing", "Water & wastewater treatment", 8, "m³", "EF21", "Scope 3 Cat 5", ""),
        ("rej", "A3 Manufacturing", "Pulper rejects to incineration", 80, "kg", "EF22", "Scope 3 Cat 5", "Plastics/metal fraction from RCP"),
    ],
    "TOTAL PCF (kg CO₂e / t board)",
)
r = PB["tot"] + 2
cell(pb, f"A{r}", "Memo: biogenic carbon stored", bold=True)
cell(pb, f"B{r}", "kg CO₂ removed & stored in 1 t board (ISO 14067 §7.2, reported separately)")
cell(pb, f"G{r}", f"=-1000*{PB['bio_c']}*44/12", fmt=N1)
cell(pb, f"A{r+1}", "PCF per kg", bold=True)
cell(pb, f"B{r+1}", "kg CO₂e / kg board (used downstream in FC & gypsum board)")
cell(pb, f"G{r+1}", f"={PB['total']}/1000", fmt="0.000", bold=True)
PB["per_kg"] = f"PCF_Board!$G${r+1}"
stage_summary(pb, PB, ["A1 Raw materials", "A2 Transport", "A3 Manufacturing"], r + 3)

# =====================================================================
# 4. PCF Folding carton (FU = 1 t folding cartons)
# =====================================================================
fc, FC = pcf_sheet(
    "PCF_FoldingCarton",
    "PCF – Printed Folding Carton (FU: 1 t, cradle-to-gate + optional distribution)",
    "Board input from own mill (linked to PCF_Board) · trim waste returned to mill as RCP (closed loop, cut-off)",
    [
        ("trim", "Trim / die-cut waste rate (of board input)", 0.12, "%", "Converting plant waste statistics"),
        ("mill_km", "Board transport mill → converting plant", 250, "km", "Internal logistics"),
        ("pb", "PCF recycled cartonboard (kg CO₂e/kg)", f"={PB['per_kg']}", "kg/kg", "Link to PCF_Board"),
    ],
    [
        ("board", "A1 Raw materials", "Recycled cartonboard (own mill)", "=1000/(1-{trim})", "kg", f"={PB['per_kg']}", "Scope 1/2/3 (internal)", "Primary data – own mill PCF"),
        ("ink", "A1 Raw materials", "Offset printing inks", 15, "kg", "EF12", "Scope 3 Cat 1", "Supplier PCF available"),
        ("varn", "A1 Raw materials", "Overprint varnish", 8, "kg", "EF13", "Scope 3 Cat 1", ""),
        ("glue", "A1 Raw materials", "Dispersion adhesive (gluing)", 5, "kg", "EF14", "Scope 3 Cat 1", ""),
        ("film", "A1 Raw materials", "Packaging – LDPE stretch film", 6, "kg", "EF15", "Scope 3 Cat 1", ""),
        ("b_tr", "A2 Transport", "Board transport mill → plant", "={board}/1000*{mill_km}", "tkm", "EF06", "Scope 3 Cat 4", "t × km"),
        ("gas", "A3 Manufacturing", "Natural gas – heating & ink drying", 80, "kWh", "EF01", "Scope 1", "Metered"),
        ("gas_up", "A3 Manufacturing", "Natural gas – upstream", "={gas}", "kWh", "EF02", "Scope 3 Cat 3", ""),
        ("el", "A3 Manufacturing", "Electricity – printing, die-cutting, gluing", 250, "kWh", "EF03", "Scope 2", "Metered"),
        ("el_up", "A3 Manufacturing", "Electricity – upstream & losses", "={el}", "kWh", "EF04", "Scope 3 Cat 3", ""),
        ("trimw", "A3 Manufacturing", "Trim waste → own mill as RCP (closed loop)", "={board}-1000", "kg", "EF08", "Scope 3 Cat 1", "Only sorting/baling allocated (cut-off)"),
    ],
    "PCF CRADLE-TO-GATE (kg CO₂e / t)",
)
r = FC["tot"] + 1
cell(fc, f"A{r}", "Optional A4", bold=True)
cell(fc, f"B{r}", "Distribution to customer (1 t × avg outbound distance)")
cell(fc, f"C{r}", f"={CI['dist_out']}", F_LINK, N1)
cell(fc, f"D{r}", "tkm")
cell(fc, f"E{r}", "EF06")
cell(fc, f"F{r}", f"={ef(f'E{r}')}", F_LINK, "0.000")
cell(fc, f"G{r}", f"=C{r}*F{r}", fmt=N1)
cell(fc, f"A{r+1}", "PCF CRADLE-TO-CUSTOMER", bold=True, fill=FILL_TOT)
cell(fc, f"G{r+1}", f"={FC['total']}+G{r}", fmt=N1, bold=True, fill=FILL_TOT)
cell(fc, f"A{r+2}", "Memo: biogenic carbon", bold=True)
cell(fc, f"B{r+2}", "kg CO₂ stored in 1 t carton")
cell(fc, f"G{r+2}", f"=-1000*PCF_Board!{PB['bio_c']}*44/12", fmt=N1)
stage_summary(fc, FC, ["A1 Raw materials", "A2 Transport", "A3 Manufacturing"], r + 4)

# =====================================================================
# 5. PCF Gypsum board (FU = 1,000 m² → reported per m²)
# =====================================================================
gb, GB = pcf_sheet(
    "PCF_GypsumBoard",
    "PCF – Gypsum Plasterboard 12.5 mm (calculated per 1,000 m², declared per m²)",
    "EN 15804 modules A1-A3 · liner paper from own recycled board mill (linked to PCF_Board)",
    [
        ("gyp", "Raw gypsum input (dihydrate) per 1,000 m²", 9000, "kg", "Recipe; water of crystallisation released in calcination"),
        ("sh_nat", "Share natural gypsum", 0.50, "%", "Procurement mix"),
        ("sh_fgd", "Share FGD gypsum", 0.30, "%", "Procurement mix – declining (coal exit)"),
        ("sh_rec", "Share recycled gypsum", "=1-C6-C7", "%", "Remainder"),
        ("g_km", "Avg gypsum transport distance", 80, "km", "Quarry / power plant / recycler"),
    ],
    [
        ("nat", "A1 Raw materials", "Natural gypsum", "={gyp}*{sh_nat}", "kg", "EF16", "Scope 3 Cat 1", ""),
        ("fgd", "A1 Raw materials", "FGD gypsum", "={gyp}*{sh_fgd}", "kg", "EF17", "Scope 3 Cat 1", ""),
        ("rec", "A1 Raw materials", "Recycled gypsum", "={gyp}*{sh_rec}", "kg", "EF18", "Scope 3 Cat 1", ""),
        ("liner", "A1 Raw materials", "Liner board (own recycled mill)", 400, "kg", f"={PB['per_kg']}", "Scope 1/2/3 (internal)", "Primary data – own mill PCF; ~200 g/m² per face"),
        ("starch", "A1 Raw materials", "Starch (core bonding)", 30, "kg", "EF09", "Scope 3 Cat 1", ""),
        ("gf", "A1 Raw materials", "Glass fibre", 10, "kg", "EF19", "Scope 3 Cat 1", ""),
        ("add", "A1 Raw materials", "Additives & foaming agent", 5, "kg", "EF20", "Scope 3 Cat 1", "Proxy – data gap"),
        ("film", "A1 Raw materials", "Packaging – LDPE film", 5, "kg", "EF15", "Scope 3 Cat 1", ""),
        ("g_tr", "A2 Transport", "Gypsum inbound transport", "={gyp}/1000*{g_km}", "tkm", "EF06", "Scope 3 Cat 4", ""),
        ("l_tr", "A2 Transport", "Liner board transport mill → plant", 120, "tkm", "EF06", "Scope 3 Cat 4", "0.4 t × 300 km"),
        ("calc", "A3 Manufacturing", "Natural gas – calcination", 2200, "kWh", "EF01", "Scope 1", "Metered"),
        ("dry", "A3 Manufacturing", "Natural gas – board drying", 4500, "kWh", "EF01", "Scope 1", "Metered – main hotspot"),
        ("gas_up", "A3 Manufacturing", "Natural gas – upstream", "={calc}+{dry}", "kWh", "EF02", "Scope 3 Cat 3", ""),
        ("el", "A3 Manufacturing", "Electricity", 600, "kWh", "EF03", "Scope 2", "Metered"),
        ("el_up", "A3 Manufacturing", "Electricity – upstream & losses", "={el}", "kWh", "EF04", "Scope 3 Cat 3", ""),
    ],
    "TOTAL (kg CO₂e / 1,000 m²)",
)
ci["B11"] = f"=B8*PCF_FoldingCarton!{FC['board']}/1000"
ci["B12"] = f"=B9*PCF_GypsumBoard!{GB['liner']}/1000"
ci["B15"] = f"=B13*PCF_Board!{PB['rcp_ratio']}"
ci["B18"] = f"=B9*PCF_GypsumBoard!{GB['gyp']}/1000"
for a in ("B11","B12","B15","B18"): ci[a].font = F_LINK
r = GB["tot"] + 1
cell(gb, f"A{r}", "Declared unit", bold=True, fill=FILL_TOT)
cell(gb, f"B{r}", "kg CO₂e per m² gypsum board", fill=FILL_TOT)
cell(gb, f"G{r}", f"={GB['total']}/1000", fmt="0.00", bold=True, fill=FILL_TOT)
GB["per_m2"] = f"PCF_GypsumBoard!$G${r}"
cell(gb, f"B{r+1}", "kg CO₂e per t gypsum board")
cell(gb, f"G{r+1}", f"=G{r}/{CI['gyp_w']}*1000", fmt=N1)
stage_summary(gb, GB, ["A1 Raw materials", "A2 Transport", "A3 Manufacturing"], r + 3)

# =====================================================================
# 6. Supplier register
# =====================================================================
sr = wb.create_sheet("Supplier_Register")
title(sr, "Supplier Register – Volumes, Supplier-Specific PCFs & PACT Data Quality (Scope 3 Cat 1)",
      "Supplier PCF is used only if provided AND DQR ≤ threshold; otherwise secondary EF. Suppliers are fictional.")
cell(sr, "A3", "DQR acceptance threshold", bold=True)
inp(sr, "D3", 3.0, "0.0")
cell(sr, "E3", "← PACT DQR scale 1 (very good) – 5 (very poor)", border=False)
SR_HDR = ["ID", "Supplier (fictional)", "Country", "Material", "Default EF ID", "Volume (t)", "Spend (EUR k)",
          "Supplier PCF (kg CO₂e/kg)", "PCF standard", "TeR", "GeR", "TiR", "Compl.", "Reliab.", "DQR",
          "Default EF", "Applied EF", "Data basis", "Emissions (t CO₂e)", "Emissions w/ default (t CO₂e)", "Δ vs default (t CO₂e)", "Next action"]
header(sr, 5, SR_HDR)
widths(sr, {"A": 6, "B": 30, "C": 7, "D": 22, "E": 8, "F": 11, "G": 10, "H": 11, "I": 14,
            "J": 5, "K": 5, "L": 5, "M": 6, "N": 6, "O": 6, "P": 8, "Q": 8, "R": 16, "S": 12, "T": 12, "U": 11, "V": 34})
SUP = [
    ("S01", "Recycling Partner North", "DE", "Recovered paper", "EF08", 160000, 19200, 0.012, "ISO 14067", 2, 1, 2, 3, 2),
    ("S02", "Recycling Partner South", "DE", "Recovered paper", "EF08", 135000, 16200, None, "", None, None, None, None, None),
    ("S03", "Municipal Paper Collection Pool", "DE", "Recovered paper", "EF08", 92000, 10100, None, "", None, None, None, None, None),
    ("S04", "RCP Trader Benelux", "NL", "Recovered paper", "EF08", 55000, 6900, 0.020, "Own method", 4, 3, 3, 4, 4),
    ("S05", "Starch Mill Central", "DE", "Native starch", "EF09", 17200, 8600, 0.62, "PACT v3", 1, 2, 1, 2, 3),
    ("S06", "Mineral Fillers Alpine", "AT", "Calcium carbonate", "EF10", 11900, 1500, 0.06, "ISO 14067", 2, 2, 3, 2, 3),
    ("S07", "ProcessChem Solutions", "DE", "Process chemicals", "EF11", 4000, 9200, None, "", None, None, None, None, None),
    ("S08", "PrintInk Europe", "DE", "Printing ink", "EF12", 1575, 7900, 2.70, "PACT v3", 2, 2, 1, 2, 2),
    ("S09", "CoatTech Varnish", "IT", "Varnish", "EF13", 840, 3400, None, "", None, None, None, None, None),
    ("S10", "BondWell Adhesives", "DE", "Adhesive", "EF14", 525, 1600, 1.50, "ISO 14067", 2, 2, 2, 3, 3),
    ("S11", "FilmPack Poly", "PL", "Stretch film", "EF15", 855, 1700, None, "", None, None, None, None, None),
    ("S12", "Gypsum Quarry Harz", "DE", "Natural gypsum", "EF16", 202500, 4100, 0.005, "EPD EN 15804", 2, 1, 2, 2, 1),
    ("S13", "Power Plant By-Products", "DE", "FGD gypsum", "EF17", 121500, 1800, None, "", None, None, None, None, None),
    ("S14", "Gypsum Recycling West", "DE", "Recycled gypsum", "EF18", 81000, 2400, 0.007, "ISO 14067", 2, 1, 1, 2, 3),
    ("S15", "GlassFibre Components", "CZ", "Glass fibre", "EF19", 450, 900, None, "", None, None, None, None, None),
    ("S16", "BuildAdd Chemicals", "DE", "Gypsum additives", "EF20", 225, 1100, None, "", None, None, None, None, None),
]
S_FIRST = 6
for i, s in enumerate(SUP):
    r = S_FIRST + i
    for j, v in enumerate(s[:5]):
        cell(sr, f"{'ABCDE'[j]}{r}", v)
    inp(sr, f"F{r}", s[5], N0)
    inp(sr, f"G{r}", s[6], N0)
    inp(sr, f"H{r}", s[7], "0.000")
    inp(sr, f"I{r}", s[8])
    for j, col in enumerate("JKLMN"):
        inp(sr, f"{col}{r}", s[9 + j], "0")
    cell(sr, f"O{r}", f'=IF(H{r}="","",AVERAGE(J{r}:N{r}))', fmt="0.0")
    cell(sr, f"P{r}", f"={ef(f'E{r}')}", F_LINK, "0.000")
    cell(sr, f"Q{r}", f'=IF(AND(H{r}<>"",O{r}<>""),IF(O{r}<=$D$3,H{r},P{r}),P{r})', fmt="0.000")
    cell(sr, f"R{r}", f'=IF(H{r}="","Secondary (no PCF)",IF(O{r}<=$D$3,"Supplier-specific","Secondary (DQR fail)"))')
    cell(sr, f"S{r}", f"=F{r}*Q{r}", fmt=N0)
    cell(sr, f"T{r}", f"=F{r}*P{r}", fmt=N0)
    cell(sr, f"U{r}", f"=S{r}-T{r}", fmt=N0)
    cell(sr, f"V{r}", f'=IF(R{r}="Supplier-specific","Annual PCF update via PACT",IF(R{r}="Secondary (DQR fail)","Improve data quality: primary energy data","Send PCF request + survey"))', wrap=False)
S_LAST = S_FIRST + len(SUP) - 1
r = S_LAST + 1
for c in "ABCDEFGHIJKLMNOPQRSTUV":
    sr[f"{c}{r}"].fill = FILL_TOT
    sr[f"{c}{r}"].border = BOX
cell(sr, f"A{r}", "TOTAL", bold=True, fill=FILL_TOT)
for col in "FGSTU":
    cell(sr, f"{col}{r}", f"=SUM({col}{S_FIRST}:{col}{S_LAST})", fmt=N0, bold=True, fill=FILL_TOT)
cell(sr, f"O{r}", f'=IFERROR(AVERAGE(O{S_FIRST}:O{S_LAST}),"")', fmt="0.0", bold=True, fill=FILL_TOT)
S_TOT = r
dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
sr.add_data_validation(dv)
dv.add(f"J{S_FIRST}:N{S_LAST}")
sr.conditional_formatting.add(f"R{S_FIRST}:R{S_LAST}", FormulaRule(formula=[f'R{S_FIRST}="Supplier-specific"'], fill=PatternFill("solid", fgColor="C6EFCE")))
sr.conditional_formatting.add(f"R{S_FIRST}:R{S_LAST}", FormulaRule(formula=[f'R{S_FIRST}="Secondary (DQR fail)"'], fill=PatternFill("solid", fgColor="FFC7CE")))
r = S_TOT + 2
cell(sr, f"A{r}", "Legend: TeR/GeR/TiR = technological/geographical/temporal representativeness; Compl. = completeness; Reliab. = reliability (PACT Pathfinder v3). Example row S01 shows the expected input format.", border=False)
sr.freeze_panes = "C6"
SRR = lambda col: f"Supplier_Register!${col}${S_FIRST}:${col}${S_LAST}"

# =====================================================================
# 7. GHG inventory
# =====================================================================
gi = wb.create_sheet("GHG_Inventory")
title(gi, "Corporate GHG Inventory – Scope 1, 2, 3 (t CO₂e)", "GHG Protocol Corporate & Scope 3 Standard · operational control · activity × EF / 1,000")
widths(gi, {"A": 8, "B": 50, "C": 15, "D": 9, "E": 10, "F": 11, "G": 13, "H": 20, "I": 7, "J": 44})
header(gi, 4, ["Code", "Emission source", "Activity", "Unit", "EF ID", "EF (kg/unit)", "t CO₂e", "Method", "DQ (1-5)", "Comment"])
GI = {}
r = 5


def gi_line(key, code, label, act, unit, efid, method, dq, note, direct=None):
    global r
    cell(gi, f"A{r}", code)
    cell(gi, f"B{r}", label)
    cell(gi, f"C{r}", act, F_LINK if "!" in str(act) else F_BASE, N0)
    cell(gi, f"D{r}", unit)
    if direct:
        cell(gi, f"E{r}", efid)
        cell(gi, f"F{r}", f"=IFERROR(G{r}/C{r},0)", fmt="0.000")
        cell(gi, f"G{r}", direct, F_LINK, N0)
    else:
        cell(gi, f"E{r}", efid)
        cell(gi, f"F{r}", f"={ef(f'E{r}')}", F_LINK, "0.000")
        cell(gi, f"G{r}", f"=C{r}*F{r}/1000", fmt=N0)
    cell(gi, f"H{r}", method)
    inp(gi, f"I{r}", dq, "0")
    cell(gi, f"J{r}", note, wrap=False)
    GI[key] = f"GHG_Inventory!$G${r}"
    GI[key + "_act"] = f"GHG_Inventory!$C${r}"
    r += 1


def sec(text):
    global r
    section(gi, r, text, 10)
    r += 1


B, FCq, GY = CI["board"], CI["fc"], CI["gyp"]
sec("SCOPE 1 – Direct emissions")
gi_line("s1_board", "S1", "Board mills – natural gas", f"={B}*PCF_Board!{PB['gas']}", "kWh", "EF01", "Metered", 1, "Board t × kWh/t")
gi_line("s1_fc", "S1", "Folding carton plants – natural gas", f"={FCq}*PCF_FoldingCarton!{FC['gas']}", "kWh", "EF01", "Metered", 1, "")
gi_line("s1_gyp", "S1", "Gypsum board plants – natural gas (calcination + drying)", f"={GY}*(PCF_GypsumBoard!{GB['calc']}+PCF_GypsumBoard!{GB['dry']})", "kWh", "EF01", "Metered", 1, "")
gi_line("s1_diesel", "S1", "Mobile combustion – diesel", f"={CI['diesel']}", "l", "EF29", "Fuel records", 1, "")
s1_first, s1_last = 6, r - 1
sec("SCOPE 2 – Purchased electricity")
el_total = f"{B}*PCF_Board!{PB['el']}+{FCq}*PCF_FoldingCarton!{FC['el']}+{GY}*PCF_GypsumBoard!{GB['el']}"
gi_line("s2_lb", "S2-LB", "Electricity – location-based (all sites)", f"={el_total}", "kWh", "EF03", "Metered", 1, "Reported; not added to MB")
gi_line("s2_mb", "S2-MB", "Electricity – market-based (non-certified share × residual mix)", f"=C{r-1}*(1-{CI['green']})", "kWh", "EF31", "Contracts", 2, "Green share in Company_Inputs")
sec("SCOPE 3 – Value chain")
mats = [("Recovered paper", "rcp", 2), ("Native starch", "starch", 2), ("Calcium carbonate", "caco3", 2), ("Process chemicals", "chem", 4),
        ("Printing ink", "ink", 2), ("Varnish", "varn", 3), ("Adhesive", "glue", 3), ("Stretch film", "film", 3),
        ("Natural gypsum", "nat", 2), ("FGD gypsum", "fgd", 3), ("Recycled gypsum", "rec", 2), ("Glass fibre", "gf", 3), ("Gypsum additives", "add", 4)]
c1_first = r
for m, k, dq in mats:
    gi_line(f"c1_{k}", "S3-01", f"Cat 1 – {m}", f'=SUMIFS({SRR("F")},{SRR("D")},"{m}")', "t", "Register", "Hybrid (supplier/avg)", dq,
            "Supplier_Register", direct=f'=SUMIFS({SRR("S")},{SRR("D")},"{m}")')
gi_line("c1_ind", "S3-01", "Cat 1 – Indirect procurement (spend-based)", f"={CI['spend_ind']}", "EUR", "EF23", "Spend-based", 5, "Replace with supplier data over time")
c1_last = r - 1
gi_line("c2", "S3-02", "Cat 2 – Capital goods", f"={CI['capex']}", "EUR", "EF24", "Spend-based", 5, "")
gas_total = f"{GI['s1_board_act']}+{GI['s1_fc_act']}+{GI['s1_gyp_act']}"
gi_line("c3_gas", "S3-03", "Cat 3 – Natural gas upstream", f"={gas_total}", "kWh", "EF02", "Average-data", 3, "")
gi_line("c3_el", "S3-03", "Cat 3 – Electricity upstream & T&D losses", f"={GI['s2_lb_act']}", "kWh", "EF04", "Average-data", 3, "")
gi_line("c3_dsl", "S3-03", "Cat 3 – Diesel upstream", f"={CI['diesel']}", "l", "EF30", "Average-data", 3, "")
gi_line("c4_rcp", "S3-04", "Cat 4 – Inbound RCP transport", f"={CI['rcp_ext']}*PCF_Board!{PB['rcp_km']}", "tkm", "EF06", "Distance-based", 3, "External RCP t × km")
gi_line("c4_gyp", "S3-04", "Cat 4 – Inbound gypsum transport", f"={CI['gyp_raw']}*PCF_GypsumBoard!{GB['g_km']}", "tkm", "EF06", "Distance-based", 3, "")
gi_line("c4_int", "S3-04", "Cat 4 – Internal transfers board → plants (3rd-party hauliers)", f"={CI['board_fc']}*PCF_FoldingCarton!{FC['mill_km']}+{GY}*PCF_GypsumBoard!{GB['l_tr']}", "tkm", "EF06", "Distance-based", 3, "")
gi_line("c4_out", "S3-04", "Cat 4 – Outbound deliveries paid by company (FC + gypsum board)", f"=({FCq}+{GY}*{CI['gyp_w']})*{CI['dist_out']}", "tkm", "EF06", "Distance-based", 4, "GLEC: request carrier data")
gi_line("c5_water", "S3-05", "Cat 5 – Water & wastewater (board mills)", f"={B}*PCF_Board!{PB['water']}", "m³", "EF21", "Average-data", 3, "")
gi_line("c5_rej", "S3-05", "Cat 5 – Pulper rejects to incineration", f"={B}*PCF_Board!{PB['rej']}", "kg", "EF22", "Average-data", 3, "")
gi_line("c6", "S3-06", "Cat 6 – Business travel", f"={CI['travel']}", "pkm", "EF25", "Distance-based", 3, "")
gi_line("c7", "S3-07", "Cat 7 – Employee commuting", f"={CI['fte']}*2*{CI['commute_km']}*{CI['workdays']}", "pkm", "EF26", "Average-data", 4, "FTE × 2 × km × days")
gi_line("c9", "S3-09", "Cat 9 – External board collected by customers", f"={CI['board_ext']}*{CI['dist_cust']}", "tkm", "EF06", "Distance-based", 4, "")
cell(gi, f"A{r}", "S3-10"); cell(gi, f"B{r}", "Cat 10 – Processing of sold board by external converters")
cell(gi, f"C{r}", f"={CI['board_ext']}", F_LINK, N0); cell(gi, f"D{r}", "t"); cell(gi, f"E{r}", "Proxy")
cell(gi, f"F{r}", f"=(PCF_FoldingCarton!{FC['el']}*Emission_Factors!$D$7+PCF_FoldingCarton!{FC['gas']}*Emission_Factors!$D$5)/(1000/(1-PCF_FoldingCarton!{FC['trim']}))*1000", F_LINK, "0.000")
cell(gi, f"G{r}", f"=C{r}*F{r}/1000", fmt=N0); cell(gi, f"H{r}", "Proxy: own converting"); inp(gi, f"I{r}", 4, "0")
cell(gi, f"J{r}", "Own FC energy intensity per t board as proxy (kg/t)")
GI["c10"] = f"GHG_Inventory!$G${r}"; r += 1
for code, txt in [("S3-08", "Cat 8 – Upstream leased assets: not relevant (included in Scope 1/2)"),
                  ("S3-11", "Cat 11 – Use of sold products: not relevant (no direct use-phase emissions)")]:
    cell(gi, f"A{r}", code); cell(gi, f"B{r}", txt); cell(gi, f"G{r}", 0, fmt=N0); cell(gi, f"H{r}", "Excluded – justified"); r += 1
gi_line("c12_fc", "S3-12", "Cat 12 – End-of-life folding cartons & board", f"=({FCq}+{CI['board_ext']})*1000", "kg", "EF27", "Average-data", 4, "Biogenic CO₂ excluded")
gi_line("c12_gyp", "S3-12", "Cat 12 – End-of-life gypsum board", f"={GY}*{CI['gyp_w']}*1000", "kg", "EF28", "Average-data", 4, "")
for code, txt in [("S3-13", "Cat 13 – Downstream leased assets: not relevant"), ("S3-14", "Cat 14 – Franchises: not relevant"),
                  ("S3-15", "Cat 15 – Investments: not relevant (no material minority holdings)")]:
    cell(gi, f"A{r}", code); cell(gi, f"B{r}", txt); cell(gi, f"G{r}", 0, fmt=N0); cell(gi, f"H{r}", "Excluded – justified"); r += 1
GI_LAST = r - 1
r += 1
sec("SUMMARY")
tot_rows = {}
for key, label, f in [
    ("S1", "Scope 1", f'=SUMIF($A$5:$A${GI_LAST},"S1",$G$5:$G${GI_LAST})'),
    ("S2LB", "Scope 2 – location-based", f'=SUMIF($A$5:$A${GI_LAST},"S2-LB",$G$5:$G${GI_LAST})'),
    ("S2MB", "Scope 2 – market-based", f'=SUMIF($A$5:$A${GI_LAST},"S2-MB",$G$5:$G${GI_LAST})'),
    ("S3", "Scope 3", f'=SUMIF($A$5:$A${GI_LAST},"S3-*",$G$5:$G${GI_LAST})'),
    ("TOT", "TOTAL Scope 1 + 2 (market-based) + 3", None),
]:
    cell(gi, f"B{r}", label, bold=True, fill=FILL_TOT)
    if f is None:
        f = f"={GI['S1'].split('!')[1]}+{GI['S2MB'].split('!')[1]}+{GI['S3'].split('!')[1]}"
    cell(gi, f"G{r}", f, fmt=N0, bold=True, fill=FILL_TOT)
    GI[key] = f"GHG_Inventory!$G${r}"
    r += 1
cell(gi, f"B{r}", "Scope 3 share of total")
cell(gi, f"G{r}", f"=G{r-2}/G{r-1}", fmt=P1)
GI["S3share"] = f"GHG_Inventory!$G${r}"
r += 2
cell(gi, f"A{r}", "Scope 3 by category (for chart)", bold=True, border=False)
r += 1
header(gi, r, ["Code", "Category", "t CO₂e", "Share of S3"])
cat_hdr = r
cat_names = ["Purchased goods & services", "Capital goods", "Fuel- & energy-related", "Upstream transport", "Waste in operations",
             "Business travel", "Employee commuting", "Upstream leased assets", "Downstream transport", "Processing of sold products",
             "Use of sold products", "End-of-life treatment", "Downstream leased assets", "Franchises", "Investments"]
for i, n in enumerate(cat_names):
    r += 1
    code = f"S3-{i+1:02d}"
    cell(gi, f"A{r}", code)
    cell(gi, f"B{r}", f"Cat {i+1} – {n}")
    cell(gi, f"C{r}", f'=SUMIF($A$5:$A${GI_LAST},A{r},$G$5:$G${GI_LAST})', fmt=N0)
    cell(gi, f"D{r}", f"=IFERROR(C{r}/{GI['S3'].split('!')[1]},0)", fmt=P1)
cat_last = r
GI["cat_c1"] = f"GHG_Inventory!$C${cat_hdr+1}"
bar = BarChart()
bar.type = "bar"
bar.title = "Scope 3 by category (t CO₂e)"
bar.add_data(Reference(gi, min_col=3, min_row=cat_hdr, max_row=cat_last), titles_from_data=True)
bar.set_categories(Reference(gi, min_col=2, min_row=cat_hdr + 1, max_row=cat_last))
bar.legend = None
bar.height, bar.width = 10, 18
bar.y_axis.numFmt = "#,##0"
bar.x_axis.scaling.orientation = "maxMin"
bar.y_axis.delete = False
bar.x_axis.delete = False
gi.add_chart(bar, f"F{cat_hdr}")
r = cat_last + 2
cell(gi, f"A{r}", "Check", bold=True)
cell(gi, f"B{r}", "Mass balance: RCP in Supplier_Register vs required external RCP (deviation)")
cell(gi, f"G{r}", f'=IFERROR(SUMIFS({SRR("F")},{SRR("D")},"Recovered paper")/{CI["rcp_ext"]}-1,0)', fmt=P1)
cell(gi, f"H{r}", f'=IF(ABS(G{r})<=2%,"OK (≤ 2%)","CHECK volumes")')
GI["mb"] = f"GHG_Inventory!$G${r}"
r += 1
cell(gi, f"A{r}", "Check", bold=True)
cell(gi, f"B{r}", "Mass balance: gypsum in Supplier_Register vs recipe requirement (deviation)")
cell(gi, f"G{r}", f'=IFERROR((SUMIFS({SRR("F")},{SRR("D")},"Natural gypsum")+SUMIFS({SRR("F")},{SRR("D")},"FGD gypsum")+SUMIFS({SRR("F")},{SRR("D")},"Recycled gypsum"))/{CI["gyp_raw"]}-1,0)', fmt=P1)
cell(gi, f"H{r}", f'=IF(ABS(G{r})<=2%,"OK (≤ 2%)","CHECK volumes")')
GI["mb_g"] = f"GHG_Inventory!$G${r}"
gi.freeze_panes = "C5"

# =====================================================================
# 8. Supplier scorecard
# =====================================================================
sc = wb.create_sheet("Supplier_Scorecard")
title(sc, "Supplier ESG Scorecard & Risk Prioritisation", "Themes aligned with EcoVadis (Environment, Labour & Human Rights, Ethics, Sustainable Procurement) + Climate data maturity")
widths(sc, {"A": 6, "B": 30, "C": 18, "D": 10, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 6, "K": 7, "L": 7, "M": 9, "N": 9, "O": 20, "P": 11, "Q": 11, "R": 9, "S": 36})
cell(sc, "A3", "Weights", bold=True)
wlabels = [("Environment", 0.30), ("Labour & HR", 0.20), ("Ethics", 0.15), ("Sust. procurement", 0.10), ("Climate maturity", 0.25)]
for i, (l, w) in enumerate(wlabels):
    col = "EFGHI"[i]
    cell(sc, f"{col}3", l, wrap=True)
    inp(sc, f"{col}4", w, P1)
cell(sc, "J4", "=SUM(E4:I4)", fmt=P1, bold=True)
cell(sc, "K4", '=IF(ROUND(J4,4)=1,"OK","≠100%")')
header(sc, 6, ["ID", "Supplier", "Material", "EcoVadis score", "Environ.", "Labour & HR", "Ethics", "Sust. proc.", "Climate maturity",
               "SBTi (Y/N)", "ISO 14001", "ISO 50001", "PCF accepted", "Total score", "Rating", "Spend (EUR k)", "Emissions (t CO₂e)", "Risk index", "Recommended action"])
SC = [  # ecovadis, env, lab, eth, sp, sbti, 14001, 50001
    (68, 72, 65, 70, 55, "Y", "Y", "Y"), (None, 50, 55, 60, 35, "N", "Y", "N"), (None, 40, 50, 55, 30, "N", "N", "N"),
    (45, 45, 50, 45, 30, "N", "Y", "N"), (74, 78, 70, 72, 65, "Y", "Y", "Y"), (61, 65, 60, 62, 50, "N", "Y", "Y"),
    (52, 55, 55, 50, 40, "N", "Y", "N"), (70, 72, 68, 70, 60, "Y", "Y", "Y"), (None, 48, 52, 50, 40, "N", "Y", "N"),
    (58, 60, 58, 60, 45, "N", "Y", "N"), (None, 42, 45, 48, 30, "N", "N", "N"), (63, 66, 62, 60, 50, "N", "Y", "Y"),
    (None, 45, 55, 50, 35, "N", "Y", "N"), (66, 70, 62, 64, 55, "Y", "Y", "N"), (None, 50, 50, 50, 40, "N", "Y", "N"),
    (None, 45, 50, 50, 35, "N", "N", "N"),
]
C_FIRST = 7
for i, s in enumerate(SC):
    r = C_FIRST + i
    srow = S_FIRST + i
    cell(sc, f"A{r}", f"=Supplier_Register!A{srow}", F_LINK)
    cell(sc, f"B{r}", f"=Supplier_Register!B{srow}", F_LINK)
    cell(sc, f"C{r}", f"=Supplier_Register!D{srow}", F_LINK)
    inp(sc, f"D{r}", s[0], "0")
    for j, col in enumerate("EFGH"):
        inp(sc, f"{col}{r}", s[1 + j], "0")
    for j, col in enumerate("JKL"):
        inp(sc, f"{col}{r}", s[5 + j])
    cell(sc, f"M{r}", f'=IF(Supplier_Register!R{srow}="Supplier-specific","Y","N")', F_LINK)
    cell(sc, f"I{r}", f'=40*(J{r}="Y")+30*(M{r}="Y")+15*(K{r}="Y")+15*(L{r}="Y")', fmt="0")
    cell(sc, f"N{r}", f"=SUMPRODUCT(E{r}:I{r},$E$4:$I$4)", fmt="0.0")
    cell(sc, f"O{r}", f'=IF(N{r}>=70,"A – Strategic partner",IF(N{r}>=55,"B – Develop",IF(N{r}>=40,"C – Improvement plan","D – Escalate")))')
    cell(sc, f"P{r}", f"=Supplier_Register!G{srow}", F_LINK, N0)
    cell(sc, f"Q{r}", f"=Supplier_Register!S{srow}", F_LINK, N0)
    cell(sc, f"R{r}", f"=Q{r}/SUM($Q${C_FIRST}:$Q${C_FIRST+len(SC)-1})*(100-N{r})", fmt="0.0")
    cell(sc, f"S{r}", f'=IF(LEFT(O{r},1)="A","Joint decarb. roadmap, PCF exchange",IF(LEFT(O{r},1)="B","Request EcoVadis + SBTi commitment",IF(LEFT(O{r},1)="C","Corrective action plan within 6 months","Escalate to category mgmt; audit")))')
C_LAST = C_FIRST + len(SC) - 1
for col, vals in [("J", '"Y,N"'), ("K", '"Y,N"'), ("L", '"Y,N"')]:
    d = DataValidation(type="list", formula1=vals, allow_blank=True)
    sc.add_data_validation(d)
    d.add(f"{col}{C_FIRST}:{col}{C_LAST}")
for letter, color in [("A", "C6EFCE"), ("B", "FFEB9C"), ("C", "FCE4D6"), ("D", "FFC7CE")]:
    sc.conditional_formatting.add(f"O{C_FIRST}:O{C_LAST}", FormulaRule(formula=[f'LEFT(O{C_FIRST},1)="{letter}"'], fill=PatternFill("solid", fgColor=color)))
r = C_LAST + 2
cell(sc, f"A{r}", "Climate maturity = 40 pts SBTi target + 30 pts accepted PCF (from register) + 15 pts ISO 14001 + 15 pts ISO 50001. Risk index = emission share × (100 − score): higher = engage first. Thematic scores (0-100) from EcoVadis scorecard or Supplier_Survey.", border=False)
sc.freeze_panes = "C7"
SCR = lambda col: f"Supplier_Scorecard!${col}${C_FIRST}:${col}${C_LAST}"

# =====================================================================
# 9. Decarbonisation levers
# =====================================================================
dc = wb.create_sheet("Decarbonisation")
title(dc, "Decarbonisation Levers – Reduction Potential & Abatement Cost", "Baselines link to GHG_Inventory · edit blue cells · negative cost = net saving")
widths(dc, {"A": 5, "B": 52, "C": 10, "D": 13, "E": 11, "F": 13, "G": 11, "H": 9, "I": 12, "J": 13, "K": 10, "L": 16, "M": 16})
header(dc, 4, ["ID", "Lever", "Scope", "Baseline (t CO₂e)", "Lever param.", "Reduction (t CO₂e/yr)", "Capex (EUR k)", "Lifetime (yrs)", "Δ Opex (EUR k/yr)", "Abatement cost (EUR/t)", "Start", "Owner", "Priority"])
LV = [
    ("Additional renewable PPA (+50% of electricity)", "S2", f"={GI['s2_mb']}", 0.50, f"={GI['s2_lb_act']}*E{{r}}*Emission_Factors!$D$35/1000", 0, 20, 1800, "2026", "Energy procurement"),
    ("Heat pumps / MVR for board drying steam", "S1", f"={GI['s1_board']}", 0.20, "=D{r}*E{r}", 38000, 20, -900, "2027", "Mill engineering"),
    ("Electrode boiler on renewable power (board mill)", "S1", f"={GI['s1_board']}", 0.15, "=D{r}*E{r}", 9000, 15, 1500, "2028", "Mill engineering"),
    ("Heat recovery in gypsum board dryers", "S1", f"={GI['s1_gyp']}", 0.12, "=D{r}*E{r}", 6000, 15, -700, "2026", "Gypsum operations"),
    ("Biomethane for gypsum calcination (share)", "S1", f"={GI['s1_gyp']}", 0.20, f"=D{{r}}*E{{r}}*(PCF_GypsumBoard!{GB['calc']}/(PCF_GypsumBoard!{GB['calc']}+PCF_GypsumBoard!{GB['dry']}))*(1-Emission_Factors!$D$36/Emission_Factors!$D$5)", 500, 10, 1600, "2027", "Energy procurement"),
    ("Low-carbon starch (supplier engagement, −30% EF)", "S3-1", f"={GI['c1_starch']}", 0.30, "=D{r}*E{r}", 0, 10, 250, "2026", "Purchasing – chemicals"),
    ("Low-carbon / bio-based inks (−25% EF)", "S3-1", f"={GI['c1_ink']}", 0.25, "=D{r}*E{r}", 0, 10, 150, "2026", "Purchasing – packaging"),
    ("Modal shift RCP inbound to rail (share of tkm)", "S3-4", f"={GI['c4_rcp']}", 0.40, "=D{r}*E{r}*(1-Emission_Factors!$D$11/Emission_Factors!$D$10)", 0, 10, 300, "2027", "Logistics"),
    ("Board lightweighting in folding cartons (−5% grammage)", "Mixed", f"={CI['board_fc']}*PCF_Board!{PB['total']}/1000", 0.05, "=D{r}*E{r}", 1500, 10, -1200, "2026", "R&D / sales"),
    ("Electric collection trucks at RCP suppliers (−30%)", "S3-1", f"={GI['c1_rcp']}", 0.30, "=D{r}*E{r}", 0, 10, 100, "2028", "Purchasing – RCP"),
]
L_FIRST = 5
for i, lv in enumerate(LV):
    r = L_FIRST + i
    cell(dc, f"A{r}", f"L{i+1:02d}")
    cell(dc, f"B{r}", lv[0])
    cell(dc, f"C{r}", lv[1])
    cell(dc, f"D{r}", lv[2], F_LINK, N0)
    inp(dc, f"E{r}", lv[3], P1)
    cell(dc, f"F{r}", lv[4].replace("{{r}}", str(r)).replace("{r}", str(r)), fmt=N0)
    inp(dc, f"G{r}", lv[5], N0)
    inp(dc, f"H{r}", lv[6], "0")
    inp(dc, f"I{r}", lv[7], N0)
    cell(dc, f"J{r}", f'=IFERROR((G{r}/H{r}+I{r})*1000/F{r},"n/a")', fmt=N0)
    inp(dc, f"K{r}", lv[8], "@")
    inp(dc, f"L{r}", lv[9])
    cell(dc, f"M{r}", f'=IF(J{r}="n/a","n/a",IF(J{r}<0,"Quick win (saves)",IF(J{r}<=$F${L_FIRST+len(LV)+2},"High – below CO₂ price",IF(J{r}<=2*$F${L_FIRST+len(LV)+2},"Medium","Low"))))')
L_LAST = L_FIRST + len(LV) - 1
r = L_LAST + 1
for c in "ABCDEFGHIJKLM":
    dc[f"{c}{r}"].fill = FILL_TOT
    dc[f"{c}{r}"].border = BOX
cell(dc, f"B{r}", "TOTAL identified reduction potential", bold=True, fill=FILL_TOT)
cell(dc, f"F{r}", f"=SUM(F{L_FIRST}:F{L_LAST})", fmt=N0, bold=True, fill=FILL_TOT)
cell(dc, f"G{r}", f"=SUM(G{L_FIRST}:G{L_LAST})", fmt=N0, bold=True, fill=FILL_TOT)
DC_TOT = f"Decarbonisation!$F${r}"
cell(dc, f"B{r+1}", "as % of total inventory (S1 + S2 MB + S3)")
cell(dc, f"F{r+1}", f"=F{r}/{GI['TOT']}", fmt=P1)
DC_PCT = f"Decarbonisation!$F${r+1}"
cell(dc, f"B{r+2}", "Reference CO₂ price (EU ETS, EUR/t) – for priority logic", border=False)
inp(dc, f"F{r+2}", 100, N0)
cell(dc, f"B{r+4}", "Note: levers are not fully additive (L02/L03 both act on board-mill gas). Abatement cost = (capex / lifetime + Δ opex) / annual reduction; no discounting.", border=False)
bar2 = BarChart()
bar2.type = "bar"
bar2.title = "Reduction potential by lever (t CO₂e/yr)"
bar2.add_data(Reference(dc, min_col=6, min_row=4, max_row=L_LAST), titles_from_data=True)
bar2.set_categories(Reference(dc, min_col=2, min_row=L_FIRST, max_row=L_LAST))
bar2.legend = None
bar2.height, bar2.width = 9, 20
bar2.x_axis.scaling.orientation = "maxMin"
bar2.y_axis.delete = False
bar2.x_axis.delete = False
dc.add_chart(bar2, f"B{r+6}")

# =====================================================================
# 10. KPI dashboard
# =====================================================================
kp = wb.create_sheet("KPI_Dashboard", 1)
title(kp, "KPI Dashboard – Sustainable Procurement & Carbon Accounting", None)
kp["A2"] = f'="Reporting year "&{CI["year"]}&" · fictional data · all values formula-linked"'
kp["A2"].font = F_SUB
widths(kp, {"A": 6, "B": 56, "C": 15, "D": 16, "E": 13, "F": 18, "G": 44})
header(kp, 4, ["#", "KPI", "Actual", "Unit", "Target", "Status", "Definition / purpose"])
spend_all = f"SUM({SCR('P')})"
KP = [
    ("Climate", None),
    ("Total GHG emissions (S1 + S2 MB + S3)", f"={GI['TOT']}", "t CO₂e", None, "", "Corporate inventory, GHG Protocol"),
    ("Scope 1", f"={GI['S1']}", "t CO₂e", None, "", ""),
    ("Scope 2 (market-based)", f"={GI['S2MB']}", "t CO₂e", None, "", "Location-based in GHG_Inventory"),
    ("Scope 3", f"={GI['S3']}", "t CO₂e", None, "", ""),
    ("Scope 3 share of total", f"={GI['S3share']}", "%", None, "", "Shows relevance of supplier engagement"),
    ("Product Carbon Footprints", None),
    ("PCF recycled cartonboard", f"=PCF_Board!{PB['total']}", "kg CO₂e/t", 550, "lower", "Cradle-to-gate, fossil"),
    ("PCF folding carton", f"=PCF_FoldingCarton!{FC['total']}", "kg CO₂e/t", 750, "lower", "Cradle-to-gate"),
    ("PCF gypsum board", f"={GB['per_m2']}", "kg CO₂e/m²", 1.8, "lower", "A1-A3, EN 15804"),
    ("Supplier data & procurement", None),
    ("Cat 1 emissions based on supplier-specific data", f'=SUMIFS({SRR("S")},{SRR("R")},"Supplier-specific")/{S_TOT and f"Supplier_Register!$S${S_TOT}"}', "%", 0.50, "higher", "Primary data share (SBTi / PACT)"),
    ("Suppliers (by spend) with EcoVadis assessment", f'=SUMIFS({SCR("P")},{SCR("D")},">0")/{spend_all}', "%", 0.80, "higher", "Rating coverage"),
    ("Spend with A/B-rated suppliers", f'=(SUMIFS({SCR("P")},{SCR("O")},"A*")+SUMIFS({SCR("P")},{SCR("O")},"B*"))/{spend_all}', "%", 0.75, "higher", "Scorecard result"),
    ("Cat 1 emissions with suppliers having SBTi targets", f'=SUMIFS({SCR("Q")},{SCR("J")},"Y")/SUM({SCR("Q")})', "%", 0.67, "higher", "SBTi supplier engagement target"),
    ("Average supplier ESG score (spend-weighted)", f"=SUMPRODUCT({SCR('N')},{SCR('P')})/{spend_all}", "pts", 60, "higher", "0-100"),
    ("Average DQR of supplier PCFs", f"=Supplier_Register!$O${S_TOT}", "1-5", 2.5, "lower", "PACT data quality rating"),
    ("Suppliers rated D – escalation needed", f'=COUNTIF({SCR("O")},"D*")', "#", 0, "lower", ""),
    ("Circularity & data quality", None),
    ("Recycled fibre share in board", "=1", "%", 1.0, "higher", "100% RCP (incl. internal trim)"),
    ("Internal trim recycled in closed loop", f"={CI['trim']}", "t", None, "", "FC trim returned to own mill"),
    ("Recycled gypsum share", "=PCF_GypsumBoard!$C$8", "%", 0.30, "higher", "Substitutes natural & FGD gypsum"),
    ("Mass balance deviation RCP (register vs. need)", f"=ABS({GI['mb']})", "%", 0.02, "lower", "Data validation check"),
    ("Decarbonisation", None),
    ("Identified reduction potential", f"={DC_TOT}", "t CO₂e/yr", None, "", "Sum of levers"),
    ("… as % of total inventory", f"={DC_PCT}", "%", 0.20, "higher", ""),
]
r = 5
n = 0
for item in KP:
    if item[1] is None:
        section(kp, r, item[0], 7)
        r += 1
        continue
    n += 1
    label, f, unit, tgt, direction, d = item
    cell(kp, f"A{r}", n)
    cell(kp, f"B{r}", label)
    fmt = P1 if unit == "%" else ("0.00" if unit in ("kg CO₂e/m²", "1-5") else N0)
    cell(kp, f"C{r}", f, F_LINK, fmt, bold=True)
    cell(kp, f"D{r}", unit)
    if tgt is not None:
        inp(kp, f"E{r}", tgt, fmt)
        op = ">=" if direction == "higher" else "<="
        cell(kp, f"F{r}", f'=IF(C{r}{op}E{r},"✔ On track","✖ Gap")')
    else:
        cell(kp, f"E{r}", "–")
        cell(kp, f"F{r}", "Monitor")
    cell(kp, f"G{r}", d)
    r += 1
kp.conditional_formatting.add(f"F5:F{r}", FormulaRule(formula=['LEFT(F5,1)="✔"'], fill=PatternFill("solid", fgColor="C6EFCE")))
kp.conditional_formatting.add(f"F5:F{r}", FormulaRule(formula=['LEFT(F5,1)="✖"'], fill=PatternFill("solid", fgColor="FFC7CE")))
# scope split chart
r += 2
cell(kp, f"B{r}", "Emission split", bold=True, border=False)
header(kp, r + 1, ["", "Scope", "t CO₂e"], col=1)
for i, (lbl, ref) in enumerate([("Scope 1", GI["S1"]), ("Scope 2 (MB)", GI["S2MB"]), ("Scope 3", GI["S3"])]):
    cell(kp, f"B{r+2+i}", lbl)
    cell(kp, f"C{r+2+i}", f"={ref}", F_LINK, N0)
pie = PieChart()
pie.title = "Scope split (t CO₂e)"
pie.add_data(Reference(kp, min_col=3, min_row=r + 1, max_row=r + 4), titles_from_data=True)
pie.set_categories(Reference(kp, min_col=2, min_row=r + 2, max_row=r + 4))
pie.dataLabels = DataLabelList()
pie.dataLabels.showPercent = True
pie.height, pie.width = 7.5, 11
kp.add_chart(pie, f"E{r}")
kp.freeze_panes = "A5"

# =====================================================================
# 11. Supplier survey
# =====================================================================
sv = wb.create_sheet("Supplier_Survey")
title(sv, "Supplier Sustainability & CO₂ Data Survey – Template", "Send annually with PCF request; answers feed Supplier_Scorecard and Supplier_Register. Column F = sample answer.")
widths(sv, {"A": 6, "B": 22, "C": 70, "D": 18, "E": 10, "F": 26, "G": 28})
header(sv, 4, ["No.", "Section", "Question", "Answer type", "Mandatory", "Sample answer (S05)", "Feeds into"])
Q = [
    ("A General", "Legal entity, sites supplying us, contact for sustainability", "Text", "Y", "Starch Mill Central, Site X", "Register"),
    ("A General", "Do you hold a valid EcoVadis rating? Please share score & date", "Number + date", "Y", "74 / 03-2025", "Scorecard D"),
    ("B Climate", "Do you report Scope 1, 2 and 3 emissions according to the GHG Protocol?", "Y/N + file", "Y", "Y", "Climate maturity"),
    ("B Climate", "Have you set a science-based target (SBTi validated or committed)?", "Validated / Committed / No", "Y", "Validated", "Scorecard J"),
    ("B Climate", "Share of renewable electricity at supplying sites (%)", "%", "Y", "45%", "Decarbonisation"),
    ("B Climate", "Planned decarbonisation measures until 2030 incl. expected reduction", "Text", "N", "Biogas boiler 2027", "Decarbonisation"),
    ("C PCF data", "Provide the cradle-to-gate PCF of each product supplied (kg CO₂e / kg)", "Number", "Y", "0.62", "Register H"),
    ("C PCF data", "Methodology used (ISO 14067, PACT, PEF, EN 15804) and verification status", "List", "Y", "PACT v3, 3rd-party verified", "Register I"),
    ("C PCF data", "Share of primary data in the PCF (%) and reference year of data", "% + year", "Y", "80% / 2024", "Register TiR/Reliab."),
    ("C PCF data", "Secondary database & version used (e.g. ecoinvent 3.10)", "Text", "Y", "ecoinvent 3.10", "Register TeR"),
    ("C PCF data", "Allocation approach for co-products and recycled content", "Text", "Y", "Economic allocation", "Data validation"),
    ("C PCF data", "Biogenic carbon content & land-use change emissions reported separately?", "Y/N", "Y", "Y", "PCF memo items"),
    ("D Environment", "Certified environmental / energy management system (ISO 14001 / 50001 / EMAS)?", "Multi-select + certificate", "Y", "ISO 14001, ISO 50001", "Scorecard K/L"),
    ("D Environment", "Water withdrawal in water-stressed areas, effluent treatment", "Text", "N", "No water stress", "Environment score"),
    ("D Environment", "Chain-of-custody certification for fibre / bio-based materials (FSC, PEFC, ISCC)", "Multi-select", "N", "ISCC PLUS", "Environment score"),
    ("E Labour & HR", "Human rights policy and due diligence process (LkSG / CSDDD)", "Y/N + file", "Y", "Y", "Labour & HR score"),
    ("E Labour & HR", "Lost-time injury frequency rate (LTIFR) last year", "Number", "Y", "3.1", "Labour & HR score"),
    ("F Ethics", "Anti-corruption & whistleblowing system in place?", "Y/N", "Y", "Y", "Ethics score"),
    ("F Ethics", "Signed our Supplier Code of Conduct?", "Y/N", "Y", "Y", "Ethics score"),
    ("G Supply chain", "Do you assess your own suppliers on ESG criteria? Share of spend covered", "%", "N", "60%", "Sust. procurement score"),
    ("G Supply chain", "Country of origin of key raw materials (deforestation / EUDR relevance)", "Text", "Y", "EU (wheat, maize)", "Risk screening"),
]
for i, q in enumerate(Q):
    rr = 5 + i
    cell(sv, f"A{rr}", f"Q{i+1:02d}")
    for j, v in enumerate(q):
        if j == 4:
            inp(sv, f"F{rr}", v)
        else:
            cell(sv, f"{'BCDEFG'[j] if j < 4 else 'G'}{rr}", v, wrap=(j == 1))
sv.freeze_panes = "A5"

# =====================================================================
# 12. Training plan
# =====================================================================
tp = wb.create_sheet("Training_Plan")
title(tp, "Internal Training & Awareness Programme", "Target groups, formats and success metrics")
widths(tp, {"A": 6, "B": 38, "C": 26, "D": 18, "E": 12, "F": 48, "G": 30})
header(tp, 4, ["No.", "Module", "Target group", "Format", "Duration", "Key content", "Success metric"])
T = [
    ("Climate basics & GHG Protocol", "All employees", "E-learning", "45 min", "Scope 1/2/3, why CO₂ matters for customers, our targets", "≥ 90% completion"),
    ("Scope 3 & supplier data for buyers", "Purchasing / category mgmt", "Workshop", "½ day", "Cat 1 hotspots, PCF requests, PACT DQR, using Supplier_Register", "Supplier-specific share ↑"),
    ("ESG supplier assessments & EcoVadis", "Purchasing, quality", "Workshop", "3 h", "EcoVadis themes, scorecard logic, corrective action plans, LkSG/CSDDD", "Spend with A/B suppliers ↑"),
    ("PCF & LCA for sales and key account", "Sales, product mgmt", "Workshop", "3 h", "Reading PCFs, cut-off & biogenic carbon, answering customer requests", "Response time to customer requests"),
    ("Energy efficiency on the shop floor", "Mill & plant operators", "On-site toolbox talk", "30 min", "Steam & compressed-air losses, idling, dryer optimisation", "kWh per t ↓"),
    ("Green claims & communication", "Marketing, sales", "Webinar", "1 h", "EU Green Claims rules, avoiding greenwashing, recyclability claims", "No non-compliant claims"),
    ("Data quality for CO₂ KPIs", "Controlling, site energy managers", "Workshop", "2 h", "Activity data, metering, plausibility checks, mass balances, audit trail", "Zero audit findings"),
]
for i, t in enumerate(T):
    rr = 5 + i
    cell(tp, f"A{rr}", f"T{i+1:02d}")
    for j, v in enumerate(t):
        cell(tp, f"{'BCDEFG'[j]}{rr}", v, wrap=True)

wb._sheets = [wb["README"], wb["KPI_Dashboard"], wb["Company_Inputs"], wb["Emission_Factors"], wb["PCF_Board"], wb["PCF_FoldingCarton"],
              wb["PCF_GypsumBoard"], wb["Supplier_Register"], wb["GHG_Inventory"], wb["Supplier_Scorecard"], wb["Decarbonisation"],
              wb["Supplier_Survey"], wb["Training_Plan"]]
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = "1F4E3D" if ws.title in ("README", "KPI_Dashboard") else "70AD47"
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
import os
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "model", "LCA_CO2_Supplier_Toolkit.xlsx")
wb.save(out)
print("saved", out, "EF rows", 5, EF_LAST)
