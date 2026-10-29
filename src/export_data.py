"""Export the recalculated workbook values to CSV + JSON for the dashboard and charts.

Run after the workbook has been recalculated (open & save in Excel, or LibreOffice recalc).
"""
import csv
import json
import os

from openpyxl import load_workbook

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
XLSX = os.path.join(ROOT, "model", "LCA_CO2_Supplier_Toolkit.xlsx")
wb = load_workbook(XLSX, data_only=True)


def rows(ws, first, last_col, stop):
    out = []
    r = first
    while not stop(ws, r):
        out.append([ws.cell(row=r, column=c).value for c in range(1, last_col + 1)])
        r += 1
    return out


def num(v):
    return round(float(v), 4) if isinstance(v, (int, float)) else v


# --- GHG inventory -------------------------------------------------------
gi = wb["GHG_Inventory"]
inventory, summary, categories = [], {}, []
for r in range(5, gi.max_row + 1):
    a, b, c, d, e, f, g, h, i = (gi.cell(row=r, column=k).value for k in range(1, 10))
    if isinstance(a, str) and (a.startswith("S1") or a.startswith("S2") or a.startswith("S3")) and isinstance(b, str) and " – " in b and h:
        inventory.append({"code": a, "source": b, "activity": num(c), "unit": d, "ef": num(f), "t_co2e": num(g), "method": h, "dq": i})
    if b in ("Scope 1", "Scope 2 – location-based", "Scope 2 – market-based", "Scope 3", "TOTAL Scope 1 + 2 (market-based) + 3") and g is not None:
        summary[b] = num(g)
    if isinstance(a, str) and a.startswith("S3-") and isinstance(b, str) and b.startswith("Cat ") and h is None and isinstance(c, (int, float)) and d is not None and not isinstance(d, str):
        categories.append({"code": a, "category": b, "t_co2e": num(c), "share": num(d)})

# --- PCFs ----------------------------------------------------------------
pcf = {}
for sheet, key, unit in [("PCF_Board", "board", "kg CO2e / t"), ("PCF_FoldingCarton", "carton", "kg CO2e / t"),
                         ("PCF_GypsumBoard", "gypsum", "kg CO2e / 1,000 m2")]:
    ws = wb[sheet]
    items, total = [], None
    for r in range(5, ws.max_row + 1):
        a, b, g = ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value, ws.cell(row=r, column=7).value
        if isinstance(a, str) and a[:2] in ("A1", "A2", "A3") and b and isinstance(g, (int, float)):
            items.append({"stage": a, "input": b, "kg_co2e": num(g), "scope": ws.cell(row=r, column=9).value})
        if isinstance(a, str) and a.startswith(("TOTAL", "PCF CRADLE-TO-GATE")) and total is None:
            total = num(g)
    pcf[key] = {"unit": unit, "total": total, "items": items}

# --- Suppliers & scorecard -----------------------------------------------
sr, sc = wb["Supplier_Register"], wb["Supplier_Scorecard"]
suppliers = []
for i in range(16):
    r, s = 6 + i, 7 + i
    suppliers.append({
        "id": sr[f"A{r}"].value, "name": sr[f"B{r}"].value, "country": sr[f"C{r}"].value, "material": sr[f"D{r}"].value,
        "volume_t": num(sr[f"F{r}"].value), "spend_keur": num(sr[f"G{r}"].value), "supplier_pcf": num(sr[f"H{r}"].value),
        "dqr": num(sr[f"O{r}"].value) if sr[f"O{r}"].value != "" else None, "data_basis": sr[f"R{r}"].value,
        "t_co2e": num(sr[f"S{r}"].value), "ecovadis": sc[f"D{r+1}"].value, "score": num(sc[f"N{s}"].value),
        "rating": sc[f"O{s}"].value, "risk_index": num(sc[f"R{s}"].value), "sbti": sc[f"J{s}"].value, "action": sc[f"S{s}"].value,
    })

# --- Levers --------------------------------------------------------------
dc = wb["Decarbonisation"]
levers = []
for r in range(5, 15):
    levers.append({"id": dc[f"A{r}"].value, "lever": dc[f"B{r}"].value, "scope": dc[f"C{r}"].value,
                   "reduction_t": num(dc[f"F{r}"].value), "capex_keur": num(dc[f"G{r}"].value),
                   "cost_eur_t": num(dc[f"J{r}"].value), "start": dc[f"K{r}"].value, "priority": dc[f"M{r}"].value})

# --- KPIs ----------------------------------------------------------------
kp = wb["KPI_Dashboard"]
kpis, group = [], None
for r in range(5, 40):
    a, b, c, d, e, f = (kp.cell(row=r, column=k).value for k in range(1, 7))
    if a is None and b is None:
        continue
    if a is None and isinstance(b, str) and c is None and kp.cell(row=r, column=1).fill.fgColor.rgb in ("00E2EFDA", "FFE2EFDA"):
        group = b
        continue
    if isinstance(a, str) and c is None:
        group = a
        continue
    if isinstance(a, int):
        kpis.append({"group": group, "kpi": b, "actual": num(c), "unit": d, "target": None if e == "–" else num(e), "status": f})

data = {"year": wb["Company_Inputs"]["B6"].value, "summary": summary, "inventory": inventory, "scope3_categories": categories,
        "pcf": pcf, "suppliers": suppliers, "levers": levers, "kpis": kpis,
        "pcf_headline": {"board_kg_t": pcf["board"]["total"], "carton_kg_t": pcf["carton"]["total"],
                         "gypsum_kg_m2": round(pcf["gypsum"]["total"] / 1000, 3)}}

os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
with open(os.path.join(ROOT, "data", "dashboard_data.json"), "w", encoding="utf-8") as fh:
    json.dump(data, fh, ensure_ascii=False, indent=1)


def write_csv(name, recs):
    with open(os.path.join(ROOT, "data", name), "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(recs[0].keys()))
        w.writeheader()
        w.writerows(recs)


write_csv("ghg_inventory.csv", inventory)
write_csv("scope3_categories.csv", categories)
write_csv("suppliers.csv", suppliers)
write_csv("decarbonisation_levers.csv", levers)
write_csv("kpis.csv", kpis)
write_csv("pcf_line_items.csv", [dict(product=k, **it) for k, v in pcf.items() for it in v["items"]])
print(json.dumps({k: (len(v) if isinstance(v, list) else v) for k, v in data.items() if k != "pcf"}, ensure_ascii=False, indent=1))
print({k: v["total"] for k, v in pcf.items()})
